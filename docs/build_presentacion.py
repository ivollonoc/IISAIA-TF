"""Genera docs/presentacion-v1.pptx.

Uso: pip install python-pptx && python docs/build_presentacion.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

SALIDA = Path(__file__).parent / "presentacion-v1.pptx"

# Paleta: verde cancha + ámbar como acento
DARK = RGBColor(0x0E, 0x3B, 0x2E)
GREEN = RGBColor(0x1E, 0x7A, 0x4F)
LINE_DARK = RGBColor(0x2A, 0x5E, 0x4B)
MINT = RGBColor(0xE7, 0xF3, 0xEC)
ACCENT = RGBColor(0xF2, 0xA9, 0x3B)
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x5B, 0x67, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xD5, 0xE3, 0xDA)

HEAD = "Arial"
BODY = "Calibri"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------

def fondo(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def caja(slide, x, y, w, h, fill=None, borde=None, redondeado=True, forma=None):
    tipo = forma or (MSO_SHAPE.ROUNDED_RECTANGLE if redondeado else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
    if tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if borde is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = borde
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def _bullet(p, nivel_in=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(nivel_in)))
    pPr.set("indent", str(-Inches(nivel_in)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def texto(slide, x, y, w, h, parrafos, size=16, color=INK, font=BODY, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, espacio=6, shape=None):
    """parrafos: str | list[str | dict(text, size, color, font, bold, bullet, italic)]"""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    for i, item in enumerate(parrafos):
        d = {"text": item} if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = d.get("align", align)
        p.space_after = Pt(d.get("espacio", espacio))
        if d.get("bullet"):
            _bullet(p)
        r = p.add_run()
        r.text = d["text"]
        f = r.font
        f.size = Pt(d.get("size", size))
        f.bold = d.get("bold", bold)
        f.italic = d.get("italic", False)
        f.name = d.get("font", font)
        f.color.rgb = d.get("color", color)
    return shape


def titulo(slide, txt, color=DARK, sub=None):
    texto(slide, 0.6, 0.45, 12.1, 0.8, txt, size=34, bold=True, font=HEAD, color=color)
    if sub:
        texto(slide, 0.6, 1.15, 12.1, 0.4, sub, size=16, color=MUTED if color == DARK else MINT)


def numero(slide, x, y, n, d=0.55, fill=ACCENT, color=DARK, size=18):
    c = caja(slide, x, y, d, d, fill=fill, forma=MSO_SHAPE.OVAL)
    texto(slide, 0, 0, 0, 0, str(n), size=size, bold=True, font=HEAD, color=color,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=c)
    return c


def flecha(slide, x1, y1, x2, y2, color=GREEN, ancho=2.0, punta=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(ancho)
    if punta:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
    return c


# ---------- 1. Portada ----------
s = prs.slides.add_slide(BLANK)
fondo(s, DARK)
# Motivo: círculo central de una cancha
ring = caja(s, 8.3, 1.0, 5.6, 5.6, forma=MSO_SHAPE.OVAL, borde=LINE_DARK)
ring.line.width = Pt(3)
caja(s, 10.95, 3.65, 0.3, 0.3, fill=LINE_DARK, forma=MSO_SHAPE.OVAL)
texto(s, 0.8, 1.9, 8, 0.4, "TP FINAL · IISAIA · FIUBA", size=14, bold=True, color=ACCENT, font=HEAD)
texto(s, 0.8, 2.4, 9, 1.2, "Gestión de Socios", size=56, bold=True, color=WHITE, font=HEAD)
texto(s, 0.8, 3.6, 8.5, 0.6, "App web multi-tenant para clubes deportivos", size=24, color=MINT)
texto(s, 0.8, 4.6, 8, 0.4, "Propuesta v1 · MVP y stack", size=16, color=ACCENT)
s.notes_slide.notes_text_frame.text = (
    "Presentamos la propuesta v1: qué problema resolvemos, para quién, "
    "cuál es el MVP y con qué stack lo construimos."
)

# ---------- 2. Problema ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "El problema")
texto(s, 0.6, 1.7, 5.6, 4.6, [
    {"text": "Los clubes chicos y medianos gestionan socios, cuotas e inscripciones con planillas y WhatsApp.",
     "size": 24, "bold": True, "color": DARK, "font": HEAD, "espacio": 18},
    {"text": "Objetivo", "size": 14, "bold": True, "color": GREEN, "espacio": 4},
    {"text": "Un espacio por club para administrar padrón, cobros y actividades, "
             "y autogestión para cada socio.", "size": 18, "color": INK},
])
dolores = [
    ("Información dispersa", "Padrón en planillas, pagos en el banco, inscripciones por chat."),
    ("¿Quién está al día?", "No hay una vista simple de las cuotas pendientes del mes."),
    ("Cupos sin control", "Las actividades se sobrecargan o quedan lugares sin usar."),
]
for i, (h, d) in enumerate(dolores):
    y = 1.7 + i * 1.6
    caja(s, 6.9, y, 5.85, 1.4, fill=MINT)
    numero(s, 7.15, y + 0.42, i + 1)
    texto(s, 7.95, y + 0.25, 4.6, 1.0, [
        {"text": h, "size": 18, "bold": True, "color": DARK, "font": HEAD, "espacio": 4},
        {"text": d, "size": 14, "color": MUTED},
    ])
s.notes_slide.notes_text_frame.text = (
    "Hoy la información está repartida entre planillas, homebanking y chats. "
    "El club no sabe rápido quién debe la cuota y los cupos de actividades se manejan a mano."
)

# ---------- 3. Usuarios ----------
s = prs.slides.add_slide(BLANK)
fondo(s, MINT)
titulo(s, "Tres tipos de usuario", sub="Cada club tiene su propio espacio; los datos nunca se mezclan entre clubes")
usuarios = [
    ("G", "Admin general", "Equipo de la plataforma", "/admin",
     ["Da de alta clubes", "Asigna el administrador de cada club"]),
    ("C", "Admin de club", "Secretaría / comisión directiva", "/c/mi-club/admin",
     ["Gestiona socios y tipos de socio", "Registra cobros", "Arma actividades con cupo"]),
    ("S", "Socio", "Miembro del club", "/c/mi-club/socio",
     ["Consulta su cuota y pagos", "Se inscribe a actividades"]),
]
for i, (letra, nombre, quien, ruta, items) in enumerate(usuarios):
    x = 0.6 + i * 4.15
    caja(s, x, 1.85, 3.85, 4.0, fill=WHITE, borde=BORDER)
    numero(s, x + 0.35, 2.15, letra, d=0.75, fill=DARK, color=ACCENT, size=22)
    texto(s, x + 0.35, 3.1, 3.2, 1.0, [
        {"text": nombre, "size": 22, "bold": True, "color": DARK, "font": HEAD, "espacio": 2},
        {"text": quien, "size": 14, "color": MUTED},
    ])
    texto(s, x + 0.35, 4.05, 3.2, 0.3, ruta, size=12, font=MONO, color=GREEN, bold=True)
    texto(s, x + 0.35, 4.55, 3.2, 1.9,
          [{"text": t, "bullet": True, "size": 15} for t in items], espacio=6)
s.notes_slide.notes_text_frame.text = (
    "Tres roles con tres vistas. El admin general administra la plataforma, "
    "el admin de club gestiona su institución y el socio se autogestiona."
)

# ---------- 4. MVP ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "MVP v1: lo mínimo, de punta a punta", sub="Un recorrido completo por cada rol antes de sumar funcionalidades")
filas = [
    ("Admin general", "Crear club (nombre + slug) y asignar su administrador por email"),
    ("Admin de club", "Tipos de socio y cuota · alta, edición y baja de socios · "
                      "registro de pagos · actividades con cupo · resumen del mes"),
    ("Socio", "Login · estado de cuota e historial de pagos · "
              "inscribirse o cancelar actividades respetando el cupo"),
]
for i, (rol, funcs) in enumerate(filas):
    y = 1.85 + i * 1.5
    caja(s, 0.6, y, 8.3, 1.3, fill=MINT)
    pill = caja(s, 0.8, y + 0.37, 2.1, 0.56, fill=DARK)
    texto(s, 0, 0, 0, 0, rol, size=14, bold=True, color=WHITE, font=HEAD,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=pill)
    texto(s, 3.15, y + 0.15, 5.55, 1.0, funcs, size=16, anchor=MSO_ANCHOR.MIDDLE)
caja(s, 9.25, 1.85, 3.5, 4.3, fill=DARK)
texto(s, 9.55, 2.1, 3.0, 3.9, [
    {"text": "Fuera de v1", "size": 20, "bold": True, "color": ACCENT, "font": HEAD, "espacio": 12},
    {"text": "Pago online (Mercado Pago)", "bullet": True},
    {"text": "Subdominio por club", "bullet": True},
    {"text": "Calendario de eventos", "bullet": True},
    {"text": "Notificaciones", "bullet": True},
    {"text": "Reportes", "bullet": True},
], size=16, color=WHITE, espacio=8)
s.notes_slide.notes_text_frame.text = (
    "El MVP cubre los tres roles con lo mínimo para que el flujo funcione completo. "
    "Lo que queda afuera está explícito y pasa al roadmap."
)

# ---------- 5. Decisiones de alcance ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Dos recortes para llegar a una v1 funcional")
recortes = [
    ("Multi-tenant por ruta", "club.gestiondesocios.com", "gestiondesocios.com/c/club",
     "Evita DNS wildcard y dominio propio. La base ya es multi-tenant: pasar a subdominios es solo cambiar el ruteo."),
    ("Cobros manuales", "Pago online con pasarela", "El admin registra efectivo o transferencia",
     "El modelo Pago ya guarda período, monto y medio. La pasarela se conecta en v2 sin cambiar datos."),
]
for i, (h, antes, ahora, por_que) in enumerate(recortes):
    x = 0.6 + i * 6.2
    caja(s, x, 1.6, 5.9, 4.5, fill=MINT)
    texto(s, x + 0.4, 1.85, 5.1, 0.5, h, size=22, bold=True, color=DARK, font=HEAD)
    texto(s, x + 0.4, 2.55, 5.1, 0.3, "IDEA ORIGINAL", size=11, bold=True, color=MUTED)
    b1 = caja(s, x + 0.4, 2.85, 5.1, 0.6, fill=WHITE, borde=BORDER)
    texto(s, 0, 0, 0, 0, antes, size=15, color=MUTED, anchor=MSO_ANCHOR.MIDDLE,
          align=PP_ALIGN.CENTER, shape=b1)
    flecha(s, x + 2.95, 3.5, x + 2.95, 3.95)
    texto(s, x + 0.4, 4.0, 5.1, 0.3, "EN LA V1", size=11, bold=True, color=GREEN)
    b2 = caja(s, x + 0.4, 4.3, 5.1, 0.6, fill=DARK)
    texto(s, 0, 0, 0, 0, ahora, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE,
          align=PP_ALIGN.CENTER, shape=b2)
    texto(s, x + 0.4, 5.15, 5.1, 1.4, por_que, size=15, color=INK)
s.notes_slide.notes_text_frame.text = (
    "Recortamos dos cosas que suman mucha complejidad y poco valor para la demo: "
    "subdominios y pagos online. Ambas están preparadas en el diseño para v2."
)

# ---------- 6. Stack y arquitectura ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Stack y arquitectura", sub="Un solo proyecto Next.js: interfaz, servidor y API en un único deploy")
nodos = [
    (0.6, 2.4, "Navegador", "React · Tailwind"),
    (3.55, 3.3, "Next.js 15", "Server Components\nServer Actions · API REST"),
    (7.4, 2.3, "Prisma ORM", "Esquema tipado\nTransacciones"),
    (10.25, 2.5, "PostgreSQL", "Multi-tenant\npor clubId"),
]
for x, w, h, d in nodos:
    fill = DARK if h == "Next.js 15" else MINT
    c1 = WHITE if fill == DARK else DARK
    c2 = MINT if fill == DARK else MUTED
    caja(s, x, 1.85, w, 1.5, fill=fill)
    texto(s, x + 0.2, 2.0, w - 0.4, 1.25, [
        {"text": h, "size": 18, "bold": True, "color": c1, "font": HEAD, "espacio": 4},
        {"text": d, "size": 13, "color": c2},
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for x1, x2 in [(3.0, 3.55), (6.85, 7.4), (9.7, 10.25)]:
    flecha(s, x1 + 0.05, 2.6, x2 - 0.05, 2.6)
flecha(s, 5.2, 3.4, 5.2, 3.8)
auth = caja(s, 3.55, 3.85, 3.3, 0.7, fill=WHITE, borde=ACCENT)
texto(s, 0, 0, 0, 0, "NextAuth · Google + JWT (rol y club)", size=13, bold=True, color=DARK,
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=auth)
capas = [
    ("Interfaz", "Next.js · React\nTypeScript · Tailwind"),
    ("Servidor", "Route Handlers y Server Actions (Node) · Zod"),
    ("Datos", "PostgreSQL · Prisma\nNeon en producción"),
    ("Infra y calidad", "Docker · Codespaces\nVercel · GitHub Actions"),
]
for i, (h, d) in enumerate(capas):
    x = 0.6 + i * 3.08
    caja(s, x, 4.95, 2.85, 1.45, fill=WHITE, borde=BORDER)
    texto(s, x + 0.25, 5.15, 2.4, 1.4, [
        {"text": h.upper(), "size": 11, "bold": True, "color": GREEN, "espacio": 6},
        {"text": d, "size": 14, "color": INK},
    ])
s.notes_slide.notes_text_frame.text = (
    "Usamos Next.js como full-stack: la UI y el servidor viven en el mismo proyecto. "
    "Elegimos Postgres porque el dominio es relacional y necesitamos transacciones para no superar cupos. "
    "NextAuth guarda rol y club en el JWT."
)

# ---------- 7. Modelo de datos ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Modelo de datos multi-tenant")
club = caja(s, 0.6, 1.5, 2.4, 3.4, fill=DARK)
texto(s, 0.85, 1.75, 1.95, 3.0, [
    {"text": "Club", "size": 22, "bold": True, "color": WHITE, "font": HEAD, "espacio": 4},
    {"text": "id · nombre · slug", "size": 12, "color": MINT, "font": MONO, "espacio": 24},
    {"text": "Todas las tablas llevan clubId y toda consulta filtra por él.", "size": 14, "color": ACCENT},
])
entidades = {
    "Usuario": (3.5, 1.5, ["email (único)", "rol", "clubId"]),
    "Socio": (6.65, 1.5, ["nroSocio · dni", "nombre · apellido", "tipoSocioId · activo"]),
    "TipoSocio": (9.9, 1.5, ["nombre", "cuotaMensual"]),
    "Pago": (3.5, 3.6, ["periodo YYYY-MM", "monto · medio", "único(socio, periodo)"]),
    "Inscripcion": (6.65, 3.6, ["actividadId · socioId", "único(actividad, socio)"]),
    "Actividad": (9.9, 3.6, ["diaSemana · hora", "cupo · activa"]),
}
W, H = 2.85, 1.3
for nombre, (x, y, campos) in entidades.items():
    caja(s, x, y, W, H, fill=MINT, borde=BORDER)
    texto(s, x + 0.2, y + 0.15, W - 0.4, H - 0.3, [
        {"text": nombre, "size": 16, "bold": True, "color": DARK, "font": HEAD, "espacio": 6},
        *[{"text": c, "size": 11, "font": MONO, "color": INK, "espacio": 2} for c in campos],
    ])
relaciones = [
    ((3.5 + W, 2.15), (6.65, 2.15)),   # Usuario 1-1 Socio
    ((9.9, 2.15), (6.65 + W, 2.15)),   # TipoSocio 1-N Socio
    ((6.9, 1.5 + H), (5.6, 3.6)),    # Socio 1-N Pago
    ((8.1, 1.5 + H), (8.1, 3.6)),    # Socio 1-N Inscripcion
    ((9.9, 4.25), (6.65 + W, 4.25)),   # Actividad 1-N Inscripcion
]
for (x1, y1), (x2, y2) in relaciones:
    flecha(s, x1, y1, x2, y2, color=MUTED, ancho=1.5, punta=False)
texto(s, 3.5, 5.3, 9.3, 0.5,
      [{"text": "Al día = existe un Pago del período actual (se calcula, no se guarda).",
        "size": 15, "italic": True, "color": GREEN}])
s.notes_slide.notes_text_frame.text = (
    "Siete tablas. El aislamiento entre clubes se resuelve con clubId en cada tabla. "
    "Las restricciones únicas evitan pagar dos veces el mismo mes o inscribirse dos veces."
)

# ---------- 8. Demo y verificación ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Demo v1 y verificación")
pasos = [
    "El admin general crea un club y asigna su administrador",
    "El admin del club carga tipo de socio, socio y pago del mes",
    "El tablero muestra socios al día, pendientes y recaudación",
    "El socio ingresa, ve su cuota y se inscribe hasta agotar el cupo",
]
for i, p in enumerate(pasos):
    y = 1.65 + i * 1.2
    numero(s, 0.6, y, i + 1, d=0.7, fill=DARK, color=ACCENT, size=20)
    texto(s, 1.6, y, 5.9, 0.7, p, size=18, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(pasos) - 1:
        flecha(s, 0.95, y + 0.75, 0.95, y + 1.15, color=BORDER, ancho=2, punta=False)
caja(s, 8.0, 1.65, 4.75, 4.15, fill=MINT)
texto(s, 8.35, 1.95, 4.1, 4.3, [
    {"text": "Calidad", "size": 20, "bold": True, "color": DARK, "font": HEAD, "espacio": 12},
    {"text": "Tests unitarios (Vitest): validación, permisos multi-tenant y cálculo de cuota", "bullet": True},
    {"text": "CI en GitHub Actions: typecheck, tests y build contra Postgres", "bullet": True},
    {"text": "Entorno reproducible con Docker Compose y Codespaces", "bullet": True},
    {"text": "Datos demo con seed para la presentación", "bullet": True},
], size=15, espacio=10)
s.notes_slide.notes_text_frame.text = (
    "Recorremos el flujo completo con los usuarios demo: admin general, admin del club y socio. "
    "La natación tiene cupo 2 para mostrar el bloqueo por cupo."
)

# ---------- 9. Roadmap ----------
s = prs.slides.add_slide(BLANK)
fondo(s, DARK)
titulo(s, "Roadmap v2", color=WHITE, sub="Lo que dejamos preparado en el diseño")
roadmap = [
    ("Pago online", "Mercado Pago y cuota mensual automática"),
    ("Subdominios", "club.gestiondesocios.com sobre el mismo modelo"),
    ("Calendario", "Eventos puntuales y reserva de turnos"),
    ("Avisos y reportes", "Vencimientos, morosidad y asistencia"),
]
for i, (h, d) in enumerate(roadmap):
    x = 0.6 + i * 3.08
    caja(s, x, 2.1, 2.85, 2.8, fill=LINE_DARK)
    numero(s, x + 0.3, 2.4, i + 1)
    texto(s, x + 0.3, 3.25, 2.3, 1.9, [
        {"text": h, "size": 20, "bold": True, "color": WHITE, "font": HEAD, "espacio": 8},
        {"text": d, "size": 15, "color": MINT},
    ])
texto(s, 0.6, 6.2, 12, 0.4, "github.com/ivollonoc/IISAIA-TF", size=16, color=ACCENT, font=MONO, bold=True)
s.notes_slide.notes_text_frame.text = (
    "Próximos pasos: pagos online, subdominios, calendario y notificaciones. "
    "El código y la documentación están en el repositorio."
)

prs.save(SALIDA)
print(f"OK -> {SALIDA}")
