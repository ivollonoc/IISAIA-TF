"""Genera docs/presentacion-v2.pptx.

Uso: pip install python-pptx && python docs/build_presentacion_v2.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

SALIDA = Path(__file__).parent / "presentacion-v2.pptx"

# Paleta: verde cancha + ámbar como acento
DARK = RGBColor(0x0E, 0x3B, 0x2E)
GREEN = RGBColor(0x1E, 0x7A, 0x4F)
LINE_DARK = RGBColor(0x2A, 0x5E, 0x4B)
MINT = RGBColor(0xE7, 0xF3, 0xEC)
ACCENT = RGBColor(0xF2, 0xA9, 0x3B)
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x5B, 0x67, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xD5, 0xE3, 0xDA)

HEAD = "Arial"
BODY = "Calibri"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------

def fondo(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def caja(slide, x, y, w, h, fill=None, borde=None, redondeado=True, forma=None):
    tipo = forma or (MSO_SHAPE.ROUNDED_RECTANGLE if redondeado else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
    if tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if borde is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = borde
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def _bullet(p, nivel_in=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(nivel_in)))
    pPr.set("indent", str(-Inches(nivel_in)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def texto(slide, x, y, w, h, parrafos, size=16, color=INK, font=BODY, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, espacio=6, shape=None):
    """parrafos: str | list[str | dict(text, size, color, font, bold, bullet, italic)]"""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    for i, item in enumerate(parrafos):
        d = {"text": item} if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = d.get("align", align)
        p.space_after = Pt(d.get("espacio", espacio))
        if d.get("bullet"):
            _bullet(p)
        r = p.add_run()
        r.text = d["text"]
        f = r.font
        f.size = Pt(d.get("size", size))
        f.bold = d.get("bold", bold)
        f.italic = d.get("italic", False)
        f.name = d.get("font", font)
        f.color.rgb = d.get("color", color)
    return shape


def titulo(slide, txt, color=DARK, sub=None):
    texto(slide, 0.6, 0.45, 12.1, 0.8, txt, size=34, bold=True, font=HEAD, color=color)
    if sub:
        texto(slide, 0.6, 1.15, 12.1, 0.4, sub, size=16, color=MUTED if color == DARK else MINT)


def numero(slide, x, y, n, d=0.55, fill=ACCENT, color=DARK, size=18):
    c = caja(slide, x, y, d, d, fill=fill, forma=MSO_SHAPE.OVAL)
    texto(slide, 0, 0, 0, 0, str(n), size=size, bold=True, font=HEAD, color=color,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=c)
    return c


def flecha(slide, x1, y1, x2, y2, color=GREEN, ancho=2.0, punta=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(ancho)
    if punta:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
    return c


# ---------- 1. Portada ----------
s = prs.slides.add_slide(BLANK)
fondo(s, DARK)
# Motivo: círculo central de una cancha
ring = caja(s, 8.3, 1.0, 5.6, 5.6, forma=MSO_SHAPE.OVAL, borde=LINE_DARK)
ring.line.width = Pt(3)
caja(s, 10.95, 3.65, 0.3, 0.3, fill=LINE_DARK, forma=MSO_SHAPE.OVAL)
texto(s, 0.8, 1.9, 8, 0.4, "TP FINAL · IISAIA · FIUBA", size=14, bold=True, color=ACCENT, font=HEAD)
texto(s, 0.8, 2.4, 9, 1.2, "Gestión de Socios", size=56, bold=True, color=WHITE, font=HEAD)
texto(s, 0.8, 3.6, 8.5, 0.6, "Del MVP a una plataforma completa para clubes", size=24, color=MINT)
texto(s, 0.8, 4.6, 8, 0.4, "Propuesta v2 · clubdesocios.com.ar", size=16, color=ACCENT)
s.notes_slide.notes_text_frame.text = (
    "En la v1 mostramos un MVP multi-tenant con lo mínimo por rol. "
    "Hoy presentamos la v2: qué agregamos, cómo lo construimos y una demo corta."
)

# ---------- 2. El problema (slide 2 de la v1) ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "El problema")
texto(s, 0.6, 1.7, 5.6, 4.6, [
    {"text": "Los clubes chicos y medianos gestionan socios, cuotas e inscripciones con planillas y WhatsApp.",
     "size": 24, "bold": True, "color": DARK, "font": HEAD, "espacio": 18},
    {"text": "Objetivo", "size": 14, "bold": True, "color": GREEN, "espacio": 4},
    {"text": "Un espacio por club para administrar padrón, cobros y actividades, "
             "y autogestión para cada socio.", "size": 18, "color": INK},
])
dolores = [
    ("Información dispersa", "Padrón en planillas, pagos en el banco, inscripciones por chat."),
    ("¿Quién está al día?", "No hay una vista simple de las cuotas pendientes del mes."),
    ("Cupos sin control", "Las actividades se sobrecargan o quedan lugares sin usar."),
]
for i, (h, d) in enumerate(dolores):
    y = 1.7 + i * 1.6
    caja(s, 6.9, y, 5.85, 1.4, fill=MINT)
    numero(s, 7.15, y + 0.42, i + 1)
    texto(s, 7.95, y + 0.25, 4.6, 1.0, [
        {"text": h, "size": 18, "bold": True, "color": DARK, "font": HEAD, "espacio": 4},
        {"text": d, "size": 14, "color": MUTED},
    ])
s.notes_slide.notes_text_frame.text = (
    "Recordemos el punto de partida: la información está repartida entre planillas, homebanking y chats. "
    "El club no sabe rápido quién debe la cuota y los cupos de actividades se manejan a mano."
)

# ---------- 3. Resumen del MVP v1 ----------
s = prs.slides.add_slide(BLANK)
fondo(s, MINT)
titulo(s, "Lo que construimos en el MVP v1",
       sub="Un recorrido completo por cada rol, multi-tenant por ruta: /c/<slug>")
roles_v1 = [
    ("G", "Admin general", ["Crea clubes (nombre + slug)", "Asigna el admin de cada club por email"]),
    ("C", "Admin de club", ["Tipos de socio y cuota mensual", "Alta, edición y baja de socios",
                            "Registra pagos del mes", "Actividades con cupo y resumen del mes"]),
    ("S", "Socio", ["Ve su estado de cuota e historial de pagos",
                    "Se inscribe o cancela respetando el cupo"]),
]
for i, (letra, nombre, items) in enumerate(roles_v1):
    x = 0.6 + i * 4.15
    caja(s, x, 1.75, 3.85, 2.95, fill=WHITE, borde=BORDER)
    numero(s, x + 0.3, 2.0, letra, d=0.6, fill=DARK, color=ACCENT, size=20)
    texto(s, x + 1.1, 2.08, 2.6, 0.5, nombre, size=20, bold=True, color=DARK, font=HEAD)
    texto(s, x + 0.3, 2.85, 3.3, 2.0,
          [{"text": t, "bullet": True, "size": 14} for t in items], espacio=5)
pilares_v1 = [
    ("STACK", "Next.js 15 · Prisma · PostgreSQL · NextAuth con Google"),
    ("DATOS", "clubId en cada tabla; “al día” = pago del mes actual"),
    ("RECORTES", "Ruta en vez de subdominio; pagos cargados por el club"),
    ("CALIDAD", "Vitest · CI en GitHub Actions · Docker y Codespaces"),
]
for i, (h, d) in enumerate(pilares_v1):
    x = 0.6 + i * 3.0833
    caja(s, x, 5.0, 2.9, 1.55, fill=DARK)
    texto(s, x + 0.25, 5.2, 2.4, 1.2, [
        {"text": h, "size": 12, "bold": True, "color": ACCENT, "font": HEAD, "espacio": 4},
        {"text": d, "size": 14, "color": WHITE},
    ])
s.notes_slide.notes_text_frame.text = (
    "En la v1 llegamos a un MVP de punta a punta: cada rol tenía su recorrido mínimo funcionando. "
    "Hicimos dos recortes conscientes, multi-tenant por ruta y cobros registrados a mano por el club, "
    "y dejamos el diseño preparado para crecer. Eso es lo que la v2 retoma."
)

# ---------- 4. De v1 a v2 ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "De la v1 a la v2 en nueve etapas",
       sub="Cada etapa: spec → plan → implementación, y se cierra con tests, build y demo en verde")
stats = [
    ("9", "etapas iteradas"),
    ("11", "módulos por dominio"),
    ("13", "migraciones Prisma"),
    ("349", "tests automatizados"),
]
for i, (n, label) in enumerate(stats):
    x = 0.6 + i * 3.08
    caja(s, x, 1.85, 2.85, 1.9, fill=MINT)
    texto(s, x + 0.3, 2.0, 2.25, 1.1, n, size=54, bold=True, color=DARK, font=HEAD)
    texto(s, x + 0.3, 3.1, 2.25, 0.45, label, size=15, color=MUTED)
etapas = ["Base", "Cuenta corriente", "Calendario", "Agenda", "Notificaciones",
          "Acceso", "Socios", "Reportes", "Subdominios"]
W_ETAPA, PASO = 1.27, 1.35
flecha(s, 0.6 + W_ETAPA / 2, 4.6, 0.6 + 8 * PASO + W_ETAPA / 2, 4.6, color=BORDER, ancho=2, punta=False)
for i, nombre in enumerate(etapas):
    x = 0.6 + i * PASO
    numero(s, x + (W_ETAPA - 0.5) / 2, 4.35, i, d=0.5, fill=DARK, color=ACCENT, size=16)
    texto(s, x, 5.0, W_ETAPA, 0.7, nombre, size=13, bold=True, color=DARK,
          align=PP_ALIGN.CENTER)
texto(s, 0.6, 6.15, 12.1, 0.4,
      "Misma arquitectura de la v1 (Next.js + Prisma/Postgres + NextAuth, multi-tenant por clubId), "
      "ahora modular y con historial de migraciones.",
      size=14, color=GREEN)
s.notes_slide.notes_text_frame.text = (
    "No reescribimos: iteramos sobre la v1 en nueve etapas. Cada una tuvo su spec, su plan "
    "y su implementación, y se cerró con typecheck, tests y build en verde. "
    "La etapa 0 no cambió comportamiento: migraciones, tests de integración y modularización."
)

# ---------- 5. Cuenta corriente ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Cuenta corriente y pagos", sub="De “un pago por mes” a cargos y pagos por socio")
flujo = [
    ("Cargos", "Cuota social mensual\nClase o evento con precio"),
    ("Pago", "Online (pasarela simulada)\no en ventanilla"),
    ("Cuenta del club", "Acreditaciones derivadas\nde los pagos"),
]
for i, (h, d) in enumerate(flujo):
    x = 0.6 + i * 4.15
    fill = DARK if i == 1 else MINT
    caja(s, x, 1.85, 3.85, 1.75, fill=fill)
    texto(s, x + 0.3, 1.95, 3.25, 1.55, [
        {"text": h, "size": 20, "bold": True, "color": WHITE if i == 1 else DARK, "font": HEAD, "espacio": 6},
        {"text": d, "size": 14, "color": MINT if i == 1 else MUTED},
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 2:
        flecha(s, x + 3.9, 2.72, x + 4.1, 2.72)
reglas = [
    ("Al día", "Sin cargos vencidos impagos. El vencimiento es configurable por club (día 1 a 28)."),
    ("Grupo familiar", "Un titular y sus familiares: solo el titular paga, y cada cargo se acredita "
                       "en la cuenta de su socio."),
    ("Pasarela enchufable", "Interfaz Pasarela: simulada con tarjetas de prueba; "
                            "Mercado Pago queda como esqueleto."),
]
for i, (h, d) in enumerate(reglas):
    x = 0.6 + i * 4.15
    caja(s, x, 4.0, 3.85, 2.7, fill=WHITE, borde=BORDER)
    numero(s, x + 0.3, 4.3, i + 1)
    texto(s, x + 0.3, 5.05, 3.25, 1.5, [
        {"text": h, "size": 18, "bold": True, "color": DARK, "font": HEAD, "espacio": 4},
        {"text": d, "size": 14, "color": INK},
    ])
s.notes_slide.notes_text_frame.text = (
    "El modelo de pagos de la v1 quedó chico. Ahora cada socio tiene una cuenta corriente con cargos "
    "y pagos; el club ve sus acreditaciones. Se pagan cargos completos, online con una pasarela "
    "simulada o en ventanilla. El titular de un grupo familiar paga por todos."
)

# ---------- 6. Agenda y reservas ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Agenda única con reserva por clase")
# Mini vista semanal (motivo visual de la agenda)
gx, gy, gw, gh = 0.6, 1.55, 5.9, 4.75
caja(s, gx, gy, gw, gh, fill=MINT)
dias = ["L", "M", "M", "J", "V", "S", "D"]
cw = (gw - 0.4) / 7
for i, d in enumerate(dias):
    texto(s, gx + 0.2 + i * cw, gy + 0.2, cw, 0.3, d, size=13, bold=True, color=GREEN,
          align=PP_ALIGN.CENTER)
bloques = [  # (día, desde, alto, texto, relleno)
    (0, 0.4, 0.8, "Natación", DARK), (2, 0.4, 0.8, "Natación", DARK),
    (1, 1.5, 0.8, "Yoga\nlleno", ACCENT), (3, 1.5, 0.8, "Yoga\nlleno", ACCENT),
    (4, 2.6, 0.8, "Funcional", DARK), (5, 0.9, 1.1, "Clínica", GREEN),
    (2, 2.6, 0.8, "Sin\nclase", WHITE),
]
for dia, desde, alto, txt, fill in bloques:
    b = caja(s, gx + 0.24 + dia * cw, gy + 0.65 + desde, cw - 0.08, alto, fill=fill,
             borde=BORDER if fill == WHITE else None)
    texto(s, 0, 0, 0, 0, txt, size=10, bold=True,
          color=DARK if fill in (ACCENT, WHITE) else WHITE,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
texto(s, gx + 0.25, gy + 4.15, gw - 0.5, 0.4, "Vistas semana · mes · lista, sin librería de calendario",
      size=12, color=MUTED)
items = [
    ("Actividades", "Varios horarios semanales, inicio y fin. Las clases se calculan; "
                    "solo se guardan suspensiones y reprogramaciones."),
    ("Eventos", "Puntuales, con inscripción, cupo y precio opcionales."),
    ("Reserva por clase", "Cupo por clase; clases sueltas o un horario por N semanas. "
                          "Si tiene precio, genera un cargo."),
    ("Lista de espera", "Al liberarse un lugar se avisa a toda la lista; nadie entra solo."),
]
for i, (h, d) in enumerate(items):
    y = 1.55 + i * 1.22
    numero(s, 6.95, y + 0.05, i + 1)
    texto(s, 7.75, y, 5.0, 1.15, [
        {"text": h, "size": 18, "bold": True, "color": DARK, "font": HEAD, "espacio": 2},
        {"text": d, "size": 14, "color": INK},
    ])
s.notes_slide.notes_text_frame.text = (
    "Unificamos actividades y calendario en una sola agenda. Las clases no se guardan: se calculan "
    "desde los horarios y solo persistimos las excepciones. Pasamos de inscripción mensual a reserva "
    "por clase, con cupo, lista de espera y precio por clase."
)

# ---------- 7. Notificaciones y acceso ----------
s = prs.slides.add_slide(BLANK)
fondo(s, MINT)
titulo(s, "Notificaciones y acceso sin contraseñas")
# Notificaciones
caja(s, 0.6, 1.55, 6.0, 4.75, fill=WHITE, borde=BORDER)
texto(s, 0.95, 1.8, 5.3, 0.45, "Notificaciones", size=22, bold=True, color=DARK, font=HEAD)
origen = caja(s, 0.95, 2.45, 5.3, 0.6, fill=DARK)
texto(s, 0, 0, 0, 0, "notificar(evento de dominio)", size=14, bold=True, color=WHITE, font=MONO,
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=origen)
canales = [("In-app", "siempre"), ("Mail", "Resend o simulado"), ("SMS", "simulado")]
for i, (c, d) in enumerate(canales):
    x = 0.95 + i * 1.82
    flecha(s, x + 0.83, 3.1, x + 0.83, 3.35)
    pill = caja(s, x, 3.4, 1.66, 0.95, fill=MINT)
    texto(s, 0, 0, 0, 0, [
        {"text": c, "size": 15, "bold": True, "color": DARK, "font": HEAD, "espacio": 0},
        {"text": d, "size": 11, "color": MUTED},
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=pill)
texto(s, 0.95, 4.6, 5.3, 1.6, [
    {"text": "Pago acreditado, cargos generados, clase suspendida, lugar liberado y avisos manuales.",
     "size": 14, "color": INK, "espacio": 8},
    {"text": "Mail y SMS se activan por club. Se notifica después del commit: si falla, no rompe la operación.",
     "size": 14, "color": GREEN},
])
# Acceso
caja(s, 6.95, 1.55, 5.8, 4.75, fill=DARK)
texto(s, 7.3, 1.8, 5.1, 0.45, "Acceso", size=22, bold=True, color=WHITE, font=HEAD)
pasos = [
    ("Email en el home", "Sin Google ni contraseñas"),
    ("Link mágico", "15 minutos, un solo uso, se guarda solo su hash"),
    ("Invitaciones", "Al crear un club o un socio con email (7 días)"),
]
for i, (h, d) in enumerate(pasos):
    y = 2.5 + i * 1.05
    numero(s, 7.3, y, i + 1)
    texto(s, 8.1, y - 0.03, 4.4, 0.9, [
        {"text": h, "size": 17, "bold": True, "color": WHITE, "font": HEAD, "espacio": 2},
        {"text": d, "size": 14, "color": MINT},
    ])
texto(s, 7.3, 5.65, 5.1, 0.5, [{"text": "Usuarios demo marcados en la base entran directo",
                                 "size": 14, "italic": True, "color": ACCENT}])
s.notes_slide.notes_text_frame.text = (
    "Los servicios notifican por llamada directa, sin bus de eventos. In-app siempre; mail y SMS según "
    "lo que active cada club, simulados detrás de una interfaz Canal. "
    "Sacamos Google: se ingresa con link mágico o invitación por mail."
)

# ---------- 8. Socios y reportes ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Socios en bloque, familias y reportes")
cols = [
    ("Importar CSV", [
        ("1", "Vista previa", "Cada fila válida o con su error"),
        ("2", "Confirmar", "Todo o nada, con invitación opcional"),
    ]),
    ("Grupo familiar", None),
    ("Reportes y puntajes", None),
]
for i, (h, _) in enumerate(cols):
    x = 0.6 + i * 4.15
    caja(s, x, 1.55, 3.85, 4.75, fill=MINT)
    texto(s, x + 0.3, 1.8, 3.25, 0.45, h, size=20, bold=True, color=DARK, font=HEAD)
# Columna 1: pasos del CSV
for j, (n, h, d) in enumerate(cols[0][1]):
    y = 2.55 + j * 1.35
    numero(s, 0.9, y, n)
    texto(s, 1.65, y - 0.02, 2.6, 1.1, [
        {"text": h, "size": 16, "bold": True, "color": DARK, "font": HEAD, "espacio": 2},
        {"text": d, "size": 14, "color": INK},
    ])
texto(s, 0.9, 5.3, 3.25, 0.8, "Hasta 500 filas · coma o punto y coma · parser propio",
      size=12, color=MUTED)
# Columna 2: árbol titular → familiares
tit = caja(s, 5.65, 2.5, 2.05, 0.6, fill=DARK)
texto(s, 0, 0, 0, 0, "Titular", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
      anchor=MSO_ANCHOR.MIDDLE, shape=tit)
for k, nombre in enumerate(["Familiar", "Familiar"]):
    fx = 5.05 + k * 1.75
    flecha(s, 6.68, 3.12, fx + 0.75, 3.55, color=MUTED, ancho=1.5, punta=False)
    f = caja(s, fx, 3.6, 1.5, 0.55, fill=WHITE, borde=BORDER)
    texto(s, 0, 0, 0, 0, nombre, size=13, color=DARK, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE, shape=f)
texto(s, 5.05, 4.45, 3.25, 1.7, [
    {"text": "Un nivel y mismo club.", "size": 14, "color": INK, "espacio": 6},
    {"text": "El titular ve y paga los cargos de todo el grupo; “al día” sigue siendo por socio.",
     "size": 14, "color": INK},
])
# Columna 3: reportes
reps = ["Socios activos y % al día", "Emitido vs. cobrado por mes", "Morosidad (top 10)",
        "Ocupación por actividad", "Puntaje de 0,5 a 5 estrellas"]
texto(s, 9.2, 2.45, 3.3, 2.8, [{"text": r, "bullet": True} for r in reps], size=14, espacio=8)
texto(s, 9.2, 5.3, 3.25, 0.8, "Gráficos con Tremor Raw sobre Recharts", size=12, color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "Para clubes reales hacía falta cargar el padrón de una vez: importación CSV con vista previa "
    "y confirmación todo o nada. Sumamos grupo familiar con un titular que paga, "
    "y reportes del club con KPIs, recaudación, morosidad, ocupación y puntaje de actividades."
)

# ---------- 9. Subdominios ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Un subdominio por club",
       sub="El recorte de la v1 se resolvió cambiando solo el ruteo: el modelo de datos no se tocó")
ruta = [
    ("Lo que ve el usuario", "demo.clubdesocios.com.ar/socio", MINT, DARK),
    ("Middleware", "TENANT_MODE=subdomain", DARK, WHITE),
    ("Ruta interna", "/c/demo/socio", MINT, DARK),
]
for i, (h, d, fill, c) in enumerate(ruta):
    x = 0.6 + i * 4.15
    caja(s, x, 1.95, 3.85, 1.45, fill=fill)
    texto(s, x + 0.15, 2.05, 3.55, 1.25, [
        {"text": h.upper(), "size": 11, "bold": True, "color": ACCENT if fill == DARK else GREEN, "espacio": 6},
        {"text": d, "size": 13, "bold": True, "color": c, "font": MONO},
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 2:
        flecha(s, x + 3.9, 2.67, x + 4.1, 2.67)
texto(s, 4.75, 3.5, 3.85, 0.35, [{"text": "reescribe", "size": 12, "italic": True, "color": MUTED}],
      align=PP_ALIGN.CENTER)
detalles = [
    ("Detrás de un flag", "Con TENANT_MODE=path todo sigue por /c/<slug>; los previews de Vercel también."),
    ("Sesión y login", "Cookie para todos los subdominios; en la página de un club solo entran sus usuarios."),
    ("Infraestructura", "Comodín *.clubdesocios.com.ar en Vercel; NIC.ar delega a los DNS de Vercel."),
]
for i, (h, d) in enumerate(detalles):
    x = 0.6 + i * 4.15
    caja(s, x, 4.15, 3.85, 2.1, fill=WHITE, borde=BORDER)
    texto(s, x + 0.3, 4.4, 3.25, 1.7, [
        {"text": h, "size": 18, "bold": True, "color": DARK, "font": HEAD, "espacio": 6},
        {"text": d, "size": 14, "color": INK},
    ])
s.notes_slide.notes_text_frame.text = (
    "En la v1 dijimos que pasar a subdominios era solo cambiar el ruteo, y así fue: un middleware "
    "reescribe el subdominio a la ruta interna del club. Cada club tiene además su página pública."
)

# ---------- 10. Arquitectura y calidad ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Arquitectura modular y calidad")
caja(s, 0.6, 1.55, 4.6, 4.75, fill=MINT)
texto(s, 0.9, 1.8, 4.0, 0.4, "src/modules/<dominio>", size=16, bold=True, color=DARK, font=MONO)
modulos = ["acceso", "actividades", "calendario", "clubs", "cuenta-corriente", "cuotas",
           "notificaciones", "pagos-online", "reportes", "reservas", "socios"]
for k, m in enumerate(modulos):
    col, fila = k % 2, k // 2
    chip = caja(s, 0.9 + col * 2.05, 2.4 + fila * 0.6, 1.9, 0.45, fill=WHITE, borde=BORDER)
    texto(s, 0, 0, 0, 0, m, size=12, color=DARK, font=MONO, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE, shape=chip)
caja(s, 5.5, 1.55, 3.4, 4.75, fill=DARK)
texto(s, 5.8, 1.8, 2.8, 4.3, [
    {"text": "Interfaces enchufables", "size": 18, "bold": True, "color": ACCENT, "font": HEAD, "espacio": 14},
    {"text": "Pasarela", "size": 16, "bold": True, "color": WHITE, "espacio": 2},
    {"text": "simulada · Mercado Pago", "size": 14, "color": MINT, "espacio": 14},
    {"text": "Canal", "size": 16, "bold": True, "color": WHITE, "espacio": 2},
    {"text": "in-app · mail · SMS", "size": 14, "color": MINT, "espacio": 14},
    {"text": "Lo simulado se reemplaza por un proveedor real sin cambiar el dominio.",
     "size": 14, "italic": True, "color": MINT},
])
caja(s, 9.2, 1.55, 3.55, 4.75, fill=WHITE, borde=BORDER)
texto(s, 9.5, 1.8, 3.0, 0.4, "Calidad", size=18, bold=True, color=DARK, font=HEAD)
ci = ["migrate deploy", "seed", "typecheck", "tests", "build"]
for k, paso in enumerate(ci):
    y = 2.4 + k * 0.62
    b = caja(s, 9.5, y, 2.95, 0.45, fill=MINT)
    texto(s, 0, 0, 0, 0, paso, size=13, bold=True, color=DARK, font=MONO,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
    if k < len(ci) - 1:
        flecha(s, 10.97, y + 0.46, 10.97, y + 0.61, ancho=1.5)
texto(s, 9.5, 5.55, 3.0, 0.65, "CI en GitHub Actions · integración contra Postgres real",
      size=12, color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "Reorganizamos el código por dominio. Lo externo (pagos, mails, SMS) vive detrás de interfaces, "
    "así lo simulado se cambia por un proveedor real. Pasamos de db push a migraciones y sumamos "
    "tests de integración contra Postgres real, que corren en el CI en cada PR."
)

# ---------- 11. Demo ----------
s = prs.slides.add_slide(BLANK)
fondo(s, WHITE)
titulo(s, "Demo v2", sub="Con los datos del seed, sin preparación manual")
demo = [
    "La socia entra a demo.clubdesocios.com.ar con el ingreso demo",
    "Reserva una clase de Natación; Yoga está lleno y queda en lista de espera",
    "Se inscribe a la Cena anual y paga sus cargos con una tarjeta de prueba",
    "El admin del club ve el pago, suspende una clase y abre los reportes",
]
for i, p in enumerate(demo):
    y = 1.75 + i * 1.15
    numero(s, 0.6, y, i + 1, d=0.7, fill=DARK, color=ACCENT, size=20)
    texto(s, 1.6, y, 6.1, 0.7, p, size=18, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(demo) - 1:
        flecha(s, 0.95, y + 0.75, 0.95, y + 1.1, color=BORDER, ancho=2, punta=False)
caja(s, 8.1, 1.75, 4.65, 4.15, fill=MINT)
texto(s, 8.45, 2.0, 4.0, 3.8, [
    {"text": "Usuarios demo", "size": 20, "bold": True, "color": DARK, "font": HEAD, "espacio": 18},
    {"text": "Admin general", "size": 16, "bold": True, "color": GREEN, "espacio": 0},
    {"text": "crea clubes y envía avisos a sus admins", "size": 16, "espacio": 22},
    {"text": "admin@demo.club", "size": 16, "bold": True, "color": GREEN, "font": MONO, "espacio": 0},
    {"text": "admin del Club Atlético Demo", "size": 16, "espacio": 22},
    {"text": "socio@demo.club", "size": 16, "bold": True, "color": GREEN, "font": MONO, "espacio": 0},
    {"text": "Lucía, titular de un grupo familiar", "size": 16},
])
s.notes_slide.notes_text_frame.text = (
    "Recorrido de un minuto: la socia reserva, cae en lista de espera, se inscribe a un evento pago "
    "y paga online. Del lado del club, el pago se acredita, la suspensión de la clase notifica "
    "a quienes reservaron y los reportes reflejan todo."
)

# ---------- 12. Cierre ----------
s = prs.slides.add_slide(BLANK)
fondo(s, DARK)
titulo(s, "Próximos pasos", color=WHITE, sub="Lo simulado ya tiene su interfaz: falta conectar lo real")
proximos = [
    ("Pagos reales", "Conectar Mercado Pago a la interfaz Pasarela"),
    ("Mails reales", "Resend con el dominio verificado"),
    ("Carnet digital", "QR del socio y registro de asistencia"),
    ("Automatizar", "Cargos y vencimientos por cron, push del navegador"),
]
for i, (h, d) in enumerate(proximos):
    x = 0.6 + i * 3.08
    caja(s, x, 2.1, 2.85, 2.8, fill=LINE_DARK)
    numero(s, x + 0.3, 2.4, i + 1)
    texto(s, x + 0.3, 3.25, 2.3, 1.9, [
        {"text": h, "size": 20, "bold": True, "color": WHITE, "font": HEAD, "espacio": 8},
        {"text": d, "size": 15, "color": MINT},
    ])
texto(s, 0.6, 6.2, 12, 0.4, "github.com/ivollonoc/IISAIA-TF  ·  clubdesocios.com.ar",
      size=16, color=ACCENT, font=MONO, bold=True)
s.notes_slide.notes_text_frame.text = (
    "Los pagos, mails y SMS ya están detrás de interfaces: el próximo paso es conectar proveedores "
    "reales. Quedan anotados el carnet digital con QR y la automatización de vencimientos. ¡Gracias!"
)

prs.save(SALIDA)
print(f"OK -> {SALIDA}")
